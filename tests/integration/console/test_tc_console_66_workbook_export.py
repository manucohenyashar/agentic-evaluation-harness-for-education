"""`TC-CONSOLE-66` (issue #771, FR-CONSOLE-58, RS-15) — the results workbook export: the
run's marks as an `.xlsx` file, read back and compared against the store.

| Arm | Expected |
|---|---|
| (a) a graded, finalized run | the served workbook carries one row per graded paper, row for row with the store's own `select_run_grades_with_students` answer, under the marks columns |
| (b) stored text beginning with `=` | exported as a **value**, never a formula — the cell's data type is string and the text round-trips unchanged |
| (c) a run with zero graded papers | the headings, then the teacher note — an honest empty sheet, not a refusal and not a truncated one |

The store rows are read through the cohort handle (the oracle side, never the console's
report); the workbook is parsed with `openpyxl` from the bytes the HTTP read returned.

**Isolation: rung 3 (E1)** — a real store, a real console on a real loopback socket, the
export taken over the console's own HTTP surface.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

import aeh.agg  # noqa: F401 — the full migration chain before the first store open (CLAUDE.md)
import aeh.det  # noqa: F401
import aeh.extract  # noqa: F401
import aeh.grade  # noqa: F401
import aeh.ingest  # noqa: F401
import aeh.integ  # noqa: F401
import aeh.judge  # noqa: F401
import aeh.orch  # noqa: F401
import aeh.pkg  # noqa: F401
import aeh.review  # noqa: F401
import aeh.synth  # noqa: F401
from aeh.grade.exports import _MARKS_COLUMNS  # noqa: SLF001 — the exported contract itself
from aeh.grade.workbook import EMPTY_RUN_NOTE, WORKBOOK_CONTENT_TYPE
from tests.support.console_redesign_world import _request, served_console
from tests.support.console_world import seed_scored_run

pytestmark = pytest.mark.integration

WORKBOOK_PATH = "/api/v1/results/workbook"


def _fetch_workbook(port: int, run_id: str, revision: str = "1"):
    """The workbook read over HTTP: `(status, headers, body-bytes)`."""
    return _request(port, "GET", f"{WORKBOOK_PATH}?run_id={run_id}&revision={revision}")


def _load_sheet(body: bytes):
    """The served bytes as an openpyxl sheet — the read-back the oracle asks for."""
    book = openpyxl.load_workbook(io.BytesIO(body), data_only=False)
    return book["Results"]


def _store_grade_rows(store, run_id: str) -> list[dict]:
    """The store's own rows for the run at revision 1 — the same statement the export
    reads, taken directly (the oracle side)."""
    from aeh.grade.schema import GRADE_STATEMENTS

    handle = store.cohort("c-wb")
    return [dict(row) for row in handle.query(
        GRADE_STATEMENTS["select_run_grades_with_students"], run_id=run_id, revision=1,
    )]


def test_tc_console_66a_the_workbook_carries_every_graded_row_the_store_holds(
        served_console):
    """A graded, finalized run: one sheet row per graded paper, row for row with the
    store's own answer, under the marks columns the CSV export carries."""
    _server, port, store, _data_dir = served_console
    seeded = seed_scored_run(store, cohort_id="c-wb", run_id="r-wb",
                             package_id="pkg-wb", submissions=3, finalize=True)
    store_rows = _store_grade_rows(store, seeded.run_id)
    assert len(store_rows) == 3, f"the seed left {len(store_rows)} grade row(s)"

    status, headers, body = _fetch_workbook(port, seeded.run_id)

    assert status == 200, body[:200]
    assert headers.get("Content-Type") == WORKBOOK_CONTENT_TYPE
    sheet = _load_sheet(body)
    assert [cell.value for cell in sheet[1]] == list(_MARKS_COLUMNS), (
        "the sheet's first row is not the marks columns")
    assert sheet.max_row == 1 + len(store_rows), (
        f"the sheet carries {sheet.max_row - 1} data row(s) for "
        f"{len(store_rows)} graded paper(s); the workbook must not truncate or pad")
    for index, store_row in enumerate(store_rows, start=2):
        cells = [cell.value for cell in sheet[index]]
        assert cells[0] == seeded.run_id
        assert cells[1] == 1, "the revision column is not the exported revision"
        assert cells[2] == store_row["student_ref"]
        assert cells[3] == store_row["submission_id"]
        assert cells[4] == round(float(store_row["total"]), 2), (
            f"row {index}'s mark is {cells[4]!r}; the store's total is "
            f"{store_row['total']!r}")
        # openpyxl serializes an empty-string cell as a blank one, so the read-back of
        # the exporter's `or ""` columns is None: the oracle pins empty↔empty and
        # value↔value, not the representation openpyxl chose.
        assert (cells[5] or "") == (store_row["grade"] or ""), (
            f"row {index}'s grade is {cells[5]!r}; the store's grade is "
            f"{store_row['grade']!r}")
        assert cells[6] == store_row["state"]
        assert (cells[7] or "") == (store_row["full_name"] or "")


def test_tc_console_66b_stored_text_that_begins_with_equals_stays_a_value_not_a_formula(
        served_console):
    """A student whose stored full name begins with `=`: the cell the workbook carries
    for it is a string holding the text verbatim — never a formula the spreadsheet
    computes on open (RS-15)."""
    _server, port, store, _data_dir = served_console
    seeded = seed_scored_run(store, cohort_id="c-wb", run_id="r-wb",
                             package_id="pkg-wb", submissions=2, finalize=True)
    with store.cohort(seeded.cohort_id).transaction() as tx:
        for submission_id in seeded.submissions:
            tx.execute(
                "INSERT INTO roster (cohort_id, student_ref, full_name) "
                "VALUES (:c, :r, :n)",
                c=seeded.cohort_id, r=f"ref-{submission_id}", n="=1+1|cmd",
            )

    _status, _headers, body = _fetch_workbook(port, seeded.run_id)
    sheet = _load_sheet(body)

    name_column = list(_MARKS_COLUMNS).index("full_name") + 1
    for index in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=index, column=name_column)
        assert cell.value == "=1+1|cmd", (
            f"row {index}'s name cell reads {cell.value!r}; the stored text must "
            "round-trip verbatim")
        assert cell.data_type == "s", (
            f"row {index}'s name cell is typed {cell.data_type!r}; text beginning "
            "with '=' must be forced back to a string cell, never a formula")
    assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row), (
        "the sheet carries a formula cell; the export must never emit one")


def test_tc_console_66c_a_run_with_no_graded_papers_exports_headings_and_the_teacher_note(
        served_console):
    """A run whose grading has not landed: the sheet carries the headings and the
    teacher sentence beside them — not a refusal, and not an empty sheet that reads as
    a lost run (RS-15's boundary)."""
    _server, port, store, _data_dir = served_console
    seeded = seed_scored_run(store, cohort_id="c-wb", run_id="r-wb",
                             package_id="pkg-wb", submissions=0, finalize=False)
    store_rows = _store_grade_rows(store, seeded.run_id)
    assert store_rows == [], "the zero-graded boundary needs a run with no grade rows"

    status, _headers, body = _fetch_workbook(port, seeded.run_id)

    assert status == 200, body[:200]
    sheet = _load_sheet(body)
    assert [cell.value for cell in sheet[1]] == list(_MARKS_COLUMNS)
    assert sheet.max_row == 2, (
        f"the empty run's sheet carries {sheet.max_row} row(s); the headings and the "
        "note are all it may hold")
    assert sheet.cell(row=2, column=1).value == EMPTY_RUN_NOTE, (
        f"the sheet's second row is {sheet.cell(row=2, column=1).value!r}; the "
        "teacher note must stand beside the headings")
